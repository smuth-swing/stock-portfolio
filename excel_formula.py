"""
excel_formula.py — 엑셀 수식 자동 평가 및 캐시 복구 모듈
- openpyxl로 엑셀 파일을 읽을 때 계산된 캐시값(<v>)이 누락된 수식 셀(=SUM, 사칙연산 등)을 직접 계산
- server.py 및 export_to_json.py에서 공통으로 사용
"""

import io
import re
import openpyxl
from openpyxl.utils import coordinate_to_tuple, get_column_letter


def resolve_cell_value(ws_val, ws_formula, cell_coord, visited=None):
    """
    특정 셀의 값을 반환합니다.
    캐시된 값(data_only=True)이 있으면 우선 사용하고,
    없으면서 수식이 있으면 수식을 파싱/계산하여 반환합니다.
    """
    if visited is None:
        visited = set()
    coord_upper = cell_coord.upper()
    if coord_upper in visited:
        return 0.0
    visited.add(coord_upper)

    try:
        val = ws_val[coord_upper].value
    except Exception:
        return None

    # 이미 계산된 유효한 숫자/값인 경우 (수식이 아닌 경우)
    if val is not None and not str(val).startswith('='):
        return val

    # 수식 확인
    try:
        formula_cell = ws_formula[coord_upper].value
    except Exception:
        return None

    if formula_cell is None:
        return val  # 둘 다 None이면 None 반환

    formula_str = str(formula_cell).strip()
    if not formula_str.startswith('='):
        try:
            return float(formula_str)
        except (ValueError, TypeError):
            return formula_str

    # '=' 제거 후 대문자화
    expr = formula_str[1:].strip().upper()

    # 1. SUM(A1:B10) 단독 처리
    sum_match = re.match(r'^SUM\(([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)$', expr)
    if sum_match:
        start_coord, end_coord = sum_match.groups()
        s_row, s_col = coordinate_to_tuple(start_coord)
        e_row, e_col = coordinate_to_tuple(end_coord)
        total = 0.0
        for r in range(min(s_row, e_row), max(s_row, e_row) + 1):
            for c in range(min(s_col, e_col), max(s_col, e_col) + 1):
                c_coord = f"{get_column_letter(c)}{r}"
                v = resolve_cell_value(ws_val, ws_formula, c_coord, visited.copy())
                try:
                    if v is not None and v != '':
                        total += float(v)
                except (ValueError, TypeError):
                    pass
        return total

    # 2. AVERAGE(A1:B10) 단독 처리
    avg_match = re.match(r'^AVERAGE\(([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)$', expr)
    if avg_match:
        start_coord, end_coord = avg_match.groups()
        s_row, s_col = coordinate_to_tuple(start_coord)
        e_row, e_col = coordinate_to_tuple(end_coord)
        total = 0.0
        count = 0
        for r in range(min(s_row, e_row), max(s_row, e_row) + 1):
            for c in range(min(s_col, e_col), max(s_col, e_col) + 1):
                c_coord = f"{get_column_letter(c)}{r}"
                v = resolve_cell_value(ws_val, ws_formula, c_coord, visited.copy())
                try:
                    if v is not None and v != '':
                        total += float(v)
                        count += 1
                except (ValueError, TypeError):
                    pass
        return (total / count) if count > 0 else 0.0

    # 3. 셀 참조 및 사칙연산 (+, -, *, /, 괄호)
    # 수식 내 SUM(...) 함수가 인라인으로 포함된 경우 먼저 치환
    def replace_inline_sum(m):
        st, en = m.group(1), m.group(2)
        s_r, s_c = coordinate_to_tuple(st)
        e_r, e_c = coordinate_to_tuple(en)
        tot = 0.0
        for r in range(min(s_r, e_r), max(s_r, e_r) + 1):
            for c in range(min(s_c, e_c), max(s_c, e_c) + 1):
                c_c = f"{get_column_letter(c)}{r}"
                v = resolve_cell_value(ws_val, ws_formula, c_c, visited.copy())
                try:
                    if v is not None and v != '':
                        tot += float(v)
                except (ValueError, TypeError):
                    pass
        return str(tot)

    expr = re.sub(r'SUM\(([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)', replace_inline_sum, expr)

    # 개별 셀 참조 치환 (예: D15, C15, F15, I15)
    def replace_cell_ref(m):
        c_coord = m.group(0)
        v = resolve_cell_value(ws_val, ws_formula, c_coord, visited.copy())
        try:
            if v is None or v == '':
                return '0.0'
            return str(float(v))
        except (ValueError, TypeError):
            return '0.0'

    eval_expr = re.sub(r'\b[A-Z]+[0-9]+\b', replace_cell_ref, expr)

    # 안전성 검증: 허용된 연산자와 숫자만 있는지 확인
    if re.match(r'^[0-9\.\+\-\*\/\(\)\s]+$', eval_expr):
        try:
            result = eval(eval_expr, {"__builtins__": {}}, {})
            return result
        except ZeroDivisionError:
            return 0.0
        except Exception:
            return 0.0

    return val if val is not None else 0.0


def get_resolved_sheet_grid(file_data: bytes, sheet_name: str):
    """
    주어진 엑셀 바이너리와 시트명에 대해,
    모든 셀의 수식을 평가한 2차원 리스트(행 x 열) 및 최대 행/열 정보를 반환합니다.
    반환: (grid_values, max_row, max_col)
    - grid_values[r-1][c-1] 형태로 1-based 인덱스 매핑 지원
    """
    wb_f = openpyxl.load_workbook(io.BytesIO(file_data), data_only=False)
    wb_v = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
    
    if sheet_name not in wb_f.sheetnames:
        return [], 0, 0

    ws_f = wb_f[sheet_name]
    ws_v = wb_v[sheet_name]

    max_r = max(ws_f.max_row, ws_v.max_row)
    max_c = max(ws_f.max_column, ws_v.max_column)

    grid = []
    for r in range(1, max_r + 1):
        row_vals = []
        for c in range(1, max_c + 1):
            coord = f"{get_column_letter(c)}{r}"
            val = resolve_cell_value(ws_v, ws_f, coord)
            row_vals.append(val)
        grid.append(row_vals)

    wb_f.close()
    wb_v.close()
    return grid, max_r, max_c
